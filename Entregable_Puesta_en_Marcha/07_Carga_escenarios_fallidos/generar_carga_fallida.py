"""
Genera un Excel de carga que la plataforma rechaza. Cada fila provoca un solo error
conocido de las validaciones de la carga; la hoja "Escenarios" indica la columna y el
mensaje que debe aparecer en el reporte de rechazo (PDF). Las dos primeras filas son
válidas: sirven para comprobar que un archivo con errores se rechaza completo y no se
guarda ningún registro.

Las fechas se calculan respecto a hoy y los códigos DTE son nuevos en cada ejecución,
así que el archivo no caduca ni choca con documentos ya cargados.

Uso:
    python3 generar_carga_fallida.py [ruta_salida.xlsx]
"""

import random
import string
import sys
import uuid
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Fecha Emision",
    "Monto",
    "Numero Documento",
    "Codigo Generacion",
    "Sello Recepcion",
    "Numero Control",
    "Metodo Emision",
    "Tipo Factura",
    "NIT Proveedor",
    "Nombre Proveedor",
    "Politica Pago",
    "Dia Desembolso",
    "Cuenta Bancaria",
]

MAIN_SUPPLIER = {
    "NIT Proveedor": "0614-150150-150-5",
    "Nombre Proveedor": "Proveedor Escenarios Fallidos S.A. de C.V.",
    "Politica Pago": "P30",
    "Dia Desembolso": "T_PLUS_1",
    "Cuenta Bancaria": "300400500601",
}


def stamp(length: int = 40) -> str:
    return "".join(random.choices(string.ascii_uppercase + string.digits, k=length))


def control(type_code: str) -> str:
    return f"DTE-{type_code}-M001P001-{random.randint(10**14, 10**15 - 1)}"


def digital(today: date, invoice_type: str = "CCF", **overrides) -> dict:
    row = {
        "Fecha Emision": today - timedelta(days=10),
        "Monto": 1250.00,
        "Numero Documento": f"DTE-{uuid.uuid4().hex[:10].upper()}",
        "Codigo Generacion": str(uuid.uuid4()).upper(),
        "Sello Recepcion": stamp(),
        "Numero Control": control("03" if invoice_type == "CCF" else "01"),
        "Metodo Emision": "DIGITAL",
        "Tipo Factura": invoice_type,
        **MAIN_SUPPLIER,
    }
    row.update(overrides)
    return row


def paper(today: date, **overrides) -> dict:
    row = {
        "Fecha Emision": today - timedelta(days=10),
        "Monto": 980.50,
        "Numero Documento": f"FAC-{uuid.uuid4().hex[:8].upper()}",
        "Codigo Generacion": None,
        "Sello Recepcion": None,
        "Numero Control": None,
        "Metodo Emision": "PAPER",
        "Tipo Factura": "FCI",
        **MAIN_SUPPLIER,
    }
    row.update(overrides)
    return row


def other_supplier(nit: str, name: str, account: str) -> dict:
    return {"NIT Proveedor": nit, "Nombre Proveedor": name, "Cuenta Bancaria": account,
            "Politica Pago": "P45", "Dia Desembolso": "T_PLUS_1"}


def build_scenarios(today: date):
    valid_digital = digital(today)
    valid_paper = paper(today)
    old_days = 200

    # (fila, columna del reporte, escenario, mensaje esperado)
    return [
        (valid_digital, "—", "Registro válido", "Sin error: no se guarda porque el archivo se rechaza completo."),
        (valid_paper, "—", "Registro válido (físico)", "Sin error: no se guarda porque el archivo se rechaza completo."),
        (digital(today, **{"Fecha Emision": today + timedelta(days=5)}),
         "Fecha de emision", "Fecha de emisión futura", "La fecha de emisión no puede ser futura."),
        (digital(today, **{"Fecha Emision": today - timedelta(days=old_days)}),
         "Fecha de emision", "Factura demasiado antigua",
         "La factura supera la antigüedad máxima permitida de 120 días."),
        (digital(today, **{"Fecha Emision": (today - timedelta(days=10)).strftime("%d/%m/%Y")}),
         "Fecha Emision", "Fecha escrita como texto",
         "La celda no tiene el formato correcto. Se requiere que aplique el formato de celda 'Fecha'..."),
        (digital(today, Monto=0), "Monto", "Monto en cero", "El monto nominal de la factura debe ser mayor a cero."),
        (digital(today, Monto=-150.00), "Monto", "Monto negativo",
         "El monto nominal de la factura debe ser mayor a cero."),
        (digital(today, Monto=100.125), "Monto", "Monto con 3 decimales",
         "El monto admite como máximo 2 decimales y la celda contiene 100.125..."),
        (digital(today, Monto="$1,000.00"), "Monto", "Monto escrito como texto",
         "La celda no tiene el formato correcto. Se requiere que aplique el formato de celda 'Número' o 'Contabilidad'..."),
        (digital(today, **{"Metodo Emision": "ELECTRONICO"}), "Forma de emision", "Método de emisión no permitido",
         "Valor no permitido. Debe ser DIGITAL o PAPER. Valor recibido: ELECTRONICO"),
        (digital(today, **{"Tipo Factura": "NC"}), "Tipo de documento", "Tipo de factura no permitido",
         "Valor no permitido. Debe ser CCF (Comprobante de Crédito Fiscal) o FCI (Factura de consumidor final)..."),
        (digital(today, **{"Codigo Generacion": None}), "Codigo de generacion", "DIGITAL sin código de generación",
         "El campo es obligatorio y no puede estar vacío."),
        (digital(today, **{"Codigo Generacion": "ABC-123"}), "Codigo de generacion", "Código de generación inválido",
         "El código de generación debe ser un UUID de 36 caracteres (8-4-4-4-12, hexadecimal)."),
        (digital(today, **{"Sello Recepcion": stamp(39)}), "Sello de recepcion", "Sello de recepción de 39 caracteres",
         "El sello de recepción debe tener exactamente 40 caracteres, solo letras y números."),
        (digital(today, **{"Numero Control": "DTE-03-123"}), "Numero de control", "Número de control mal formado",
         "El número de control debe tener el formato DTE-TT-XXXXXXXX-NNNNNNNNNNNNNNN..."),
        (digital(today, **{"Numero Control": control("05")}), "Numero de control", "Tipo de DTE no financiable (05)",
         "El tipo de DTE 05 no es financiable. Solo se aceptan 01 y 03."),
        (digital(today, invoice_type="FCI", **{"Numero Control": control("03")}), "Numero de control",
         "Número de control no coincide con el tipo",
         "El número de control corresponde al tipo 03 (...), pero el tipo de documento indicado es FCI (01)."),
        (paper(today, **{"Codigo Generacion": str(uuid.uuid4()).upper()}), "Codigo de generacion",
         "Factura física con código de generación",
         "Los documentos en papel (PAPER) no deben contener código de generación."),
        (paper(today, **{"Numero Documento": None}), "Numero de documento", "Factura física sin número",
         "El número de documento es obligatorio para facturas físicas (PAPER)."),
        (digital(today, **{"Politica Pago": "P15"}), "Politica de pago", "Política de pago inexistente",
         "La política de pago no existe en el catálogo. Valor recibido: P15"),
        (digital(today, **{"Dia Desembolso": "LUNES"}), "Dia de desembolso", "Día de desembolso inexistente",
         "El día de desembolso no existe en el catálogo de políticas de desembolso. Valor recibido: LUNES"),
        (digital(today, **{"Nombre Proveedor": None}), "Nombre del proveedor", "Proveedor sin nombre",
         "El campo es obligatorio y no puede estar vacío."),
        (digital(today, **{"Codigo Generacion": valid_digital["Codigo Generacion"]}), "Codigo de generacion",
         "Código de generación repetido en el archivo",
         "Este código de generación viene duplicado dentro de este mismo archivo Excel."),
        (paper(today, **{"Numero Documento": valid_paper["Numero Documento"]}), "Numero de documento",
         "Factura física repetida en el archivo",
         "Esta factura física viene duplicada dentro de este mismo archivo Excel para este proveedor."),
        (digital(today, **other_supplier("0614-123", "Proveedor NIT Corto S.A.", "300400500602")),
         "NIT del proveedor", "NIT con menos de 14 dígitos",
         "El NIT debe tener exactamente 14 dígitos. Valor recibido: 0614123"),
        (digital(today, **other_supplier(6140303031033, "Proveedor NIT Numerico S.A.", "300400500603")),
         "NIT Proveedor", "NIT guardado como número",
         "La celda es numérica. Aplique el formato de celda 'Texto' en Excel..."),
        (digital(today, **other_supplier("0614-404040-104-4", "Proveedor Cuenta Invalida S.A.", "3004-0050A")),
         "Cuenta bancaria", "Cuenta bancaria con letras",
         "La cuenta bancaria solo puede contener números (máximo 50). Valor recibido: 30040050A"),
        (digital(today, **other_supplier("0614-505050-105-5", "Proveedor Sin Cuenta S.A.", None)),
         "Cuenta bancaria", "Cuenta bancaria vacía", "El campo es obligatorio y no puede estar vacío."),
        (digital(today, **other_supplier("0614-606060-106-6", "Proveedor Dos Cuentas S.A.", "300400500606")),
         "—", "Proveedor con dos cuentas (primera fila, válida)", "Sin error en esta fila."),
        (digital(today, **other_supplier("0614-606060-106-6", "Proveedor Dos Cuentas S.A.", "300400500699")),
         "Cuenta bancaria", "Proveedor con dos cuentas distintas",
         "Todas las facturas del proveedor (NIT: 06146060601066) deben tener la misma cuenta bancaria..."),
    ]


def write_workbook(scenarios, output: Path, today: date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="C8102E")

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row, *_ in scenarios:
        ws.append([row.get(h) for h in HEADERS])
        r = ws.max_row
        if isinstance(row["Fecha Emision"], date):
            ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        if isinstance(row["Monto"], (int, float)):
            ws.cell(row=r, column=2).number_format = "#,##0.000" if row["Monto"] == 100.125 else "#,##0.00"
        for name in ("NIT Proveedor", "Cuenta Bancaria"):
            if isinstance(row.get(name), str):
                ws.cell(row=r, column=HEADERS.index(name) + 1).number_format = "@"

    for idx, header in enumerate(HEADERS, start=1):
        width = max(len(header), *(len(str(s[0].get(header) or "")) for s in scenarios)) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 45)
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Escenarios")
    info.append([f"Generado el {today.strftime('%d/%m/%Y')}. La plataforma solo lee la hoja 'Documentos'."])
    info.append(["Resultado esperado: carga rechazada, reporte PDF de errores y correo de 'Fallo en la carga' "
                 "al pagador y a los usuarios bancarios. No se guarda ningún registro."])
    info.append([])
    info.append(["Fila Excel", "Escenario", "Columna en el reporte", "Mensaje esperado"])
    for cell in info[4]:
        cell.font = header_font
        cell.fill = header_fill
    for excel_row, (_, column, scenario, message) in enumerate(scenarios, start=2):
        info.append([excel_row, scenario, column, message])
    for col, width in zip("ABCD", (11, 46, 24, 110)):
        info.column_dimensions[col].width = width

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    today = date.today()
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "carga_escenarios_fallidos.xlsx"
    scenarios = build_scenarios(today)
    write_workbook(scenarios, output, today)
    errors = sum(1 for s in scenarios if s[1] != "—")
    print(f"Archivo: {output}")
    print(f"Filas: {len(scenarios)}  |  Filas con error: {errors}")


if __name__ == "__main__":
    main()
