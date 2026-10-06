"""
Genera un Excel de carga de documentos para pruebas, compatible con la plantilla
de excel_template_columns (migraciones V2 y V6 de Flyway) y las validaciones de
DocumentValidationServiceImpl. El proveedor se identifica por su NIT; la carga no
crea usuarios (se dan de alta en Gestión de usuarios).

Las fechas se calculan respecto a hoy para que siempre existan documentos:
  - vencidos y cercanos al vencimiento (pasarán a Inactivo cuando el proveedor
    consulte sus documentos),
  - financiables con distintas fechas de vencimiento.

Cada ejecución genera códigos de generación y números de documento nuevos, por
lo que el archivo puede cargarse varias veces sin chocar con la validación de
doble fondeo.

La carga se rechaza si el total supera el límite disponible del pagador (límite de la
línea × umbral de alerta − monto en uso). Con --total se reparten los montos para que
el archivo sume exactamente ese valor.

Uso:
    python3 generar_carga_exitosa.py [ruta_salida.xlsx] [--total 9500]
"""

import argparse
import random
import string
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Fecha Emision",
    "Monto",
    "Numero Documento",
    "Codigo Generacion",
    "Sello Recepcion",
    "Numero Control",
    "Metodo Emision",
    "Tipo Factura",
    "NIT Proveedor",
    "Nombre Proveedor",
    "Politica Pago",
    "Dia Desembolso",
    "Cuenta Bancaria",
]

POLICY_DAYS = {"P30": 30, "P45": 45, "P60": 60, "P90": 90}

# El primero es el proveedor demo del seed (ya existe con cuenta 000987654321 y convenio P60/T_PLUS_1).
SUPPLIERS = [
    {
        "nit": "0614-000000-003-2",
        "name": "Proveedor Demo El Salvador S.A. de C.V.",
        "policy": "P60",
        "disbursement": "T_PLUS_1",
        "account": "000987654321",
        # Días hasta el vencimiento respecto a hoy (negativo = ya vencido).
        "due_offsets": [-6, 0, 3, 12, 20, 28, 35, 45, 55],
    },
    {
        "nit": "0614-010101-101-1",
        "name": "Distribuidora Centroamericana S.A. de C.V.",
        "policy": "P30",
        "disbursement": "ONLY_FRIDAYS",
        "account": "100200300401",
        "due_offsets": [-10, -2, 4, 6, 9, 15, 22, 29],
    },
    {
        "nit": "0614-020202-102-2",
        "name": "Servicios Industriales Cuscatlan S.A. de C.V.",
        "policy": "P45",
        "disbursement": "T_PLUS_1",
        "account": "100200300402",
        "due_offsets": [2, 8, 14, 21, 30, 38, 44],
    },
    {
        "nit": "0614-030303-103-3",
        "name": "Tecnologias del Pacifico S.A. de C.V.",
        "policy": "P90",
        "disbursement": "T_PLUS_1",
        "account": "100200300403",
        "due_offsets": [5, 18, 33, 47, 60, 75, 88],
    },
]


def random_stamp() -> str:
    return "".join(random.choices(string.ascii_uppercase + string.digits, k=40))


def control_number(invoice_type: str, seq: int) -> str:
    type_code = "03" if invoice_type == "CCF" else "01"
    value = f"DTE-{type_code}-M001P001-{seq:015d}"
    assert len(value) == 31
    return value


def build_rows(today: date):
    run_tag = datetime.now().strftime("%m%d%H%M%S")
    seq_base = int(datetime.now().strftime("%y%m%d%H%M%S")) * 100
    rows = []
    seq = 0

    for s_index, supplier in enumerate(SUPPLIERS):
        days = POLICY_DAYS[supplier["policy"]]
        for d_index, offset in enumerate(supplier["due_offsets"]):
            seq += 1
            due = today + timedelta(days=offset)
            issue = due - timedelta(days=days)
            if issue > today or issue < today - timedelta(days=120):
                raise ValueError(f"Fecha de emisión fuera de rango para {supplier['name']} (offset {offset})")

            digital = (d_index % 3) != 2
            invoice_type = "CCF" if (d_index % 2 == 0) else "FCI"
            amount = round(random.uniform(350, 4800), 2)

            rows.append({
                "Fecha Emision": issue,
                "Monto": amount,
                "Numero Documento": f"{'DTE' if digital else 'FAC'}-{s_index + 1}{run_tag}-{d_index + 1:02d}",
                "Codigo Generacion": str(uuid.uuid4()).upper() if digital else None,
                "Sello Recepcion": random_stamp() if digital else None,
                "Numero Control": control_number(invoice_type, seq_base + seq) if digital else None,
                "Metodo Emision": "DIGITAL" if digital else "PAPER",
                "Tipo Factura": invoice_type,
                "NIT Proveedor": supplier["nit"],
                "Nombre Proveedor": supplier["name"],
                "Politica Pago": supplier["policy"],
                "Dia Desembolso": supplier["disbursement"],
                "Cuenta Bancaria": supplier["account"],
                "_due": due,
                "_offset": offset,
            })
    return rows


def scenario_label(offset: int) -> str:
    if offset < 0:
        return "Vencido: pasará a Inactivo (DUE_DATE_EXPIRED)"
    if offset <= 7:
        return "Cercano al vencimiento: probablemente pasará a Inactivo (NEAR_DUE_DATE_UNREQUESTED)"
    return "Financiable"


def write_workbook(rows, output: Path, today: date):
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="C8102E")

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row[h] for h in HEADERS])
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=2).number_format = "#,##0.00"
        # Identificadores como texto para conservar ceros a la izquierda.
        for col_name in ("Cuenta Bancaria", "NIT Proveedor"):
            ws.cell(row=r, column=HEADERS.index(col_name) + 1).number_format = "@"

    for idx, header in enumerate(HEADERS, start=1):
        width = max(len(header), *(len(str(r[header] or "")) for r in rows)) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 45)
    ws.freeze_panes = "A2"

    # Hoja informativa: el backend solo lee la primera hoja.
    info = wb.create_sheet("Escenarios")
    info.append([f"Generado el {today.strftime('%d/%m/%Y')} (fechas relativas a este día)"])
    info.append([])
    info.append(["Numero Documento", "Proveedor", "Politica", "Emision", "Vencimiento", "Dias al vencimiento", "Escenario esperado"])
    for cell in info[3]:
        cell.font = header_font
        cell.fill = header_fill
    for row in rows:
        info.append([
            row["Numero Documento"], row["Nombre Proveedor"], row["Politica Pago"],
            row["Fecha Emision"], row["_due"], row["_offset"], scenario_label(row["_offset"]),
        ])
        r = info.max_row
        info.cell(row=r, column=4).number_format = "DD/MM/YYYY"
        info.cell(row=r, column=5).number_format = "DD/MM/YYYY"
    for col, width in zip("ABCDEFG", (26, 46, 10, 12, 13, 20, 80)):
        info.column_dimensions[col].width = width

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def scale_amounts(rows, total: float):
    """Reparte `total` en proporción a los montos aleatorios; el redondeo se ajusta en la última fila."""
    cents = round(total * 100)
    weights = [r["Monto"] for r in rows]
    shares = [max(1, int(cents * w / sum(weights))) for w in weights]
    shares[-1] += cents - sum(shares)
    if shares[-1] <= 0:
        raise ValueError(f"El total {total} es muy bajo para {len(rows)} documentos")
    for row, share in zip(rows, shares):
        row["Monto"] = share / 100


def main():
    today = date.today()
    default_name = f"carga_documentos_{today.strftime('%Y%m%d')}.xlsx"
    parser = argparse.ArgumentParser(description="Genera un Excel de carga de documentos para pruebas.")
    parser.add_argument("output", nargs="?", type=Path, default=Path(__file__).parent / default_name)
    parser.add_argument("--total", type=float, help="Monto total exacto del archivo")
    args = parser.parse_args()
    output = args.output

    rows = build_rows(today)
    if args.total is not None:
        scale_amounts(rows, args.total)
    write_workbook(rows, output, today)

    total = sum(r["Monto"] for r in rows)
    print(f"Archivo: {output}")
    print(f"Documentos: {len(rows)}  |  Proveedores: {len(SUPPLIERS)}  |  Monto total: ${total:,.2f}")
    for label in ("Vencido", "Cercano", "Financiable"):
        count = sum(1 for r in rows if scenario_label(r["_offset"]).startswith(label))
        print(f"  {label}: {count}")


if __name__ == "__main__":
    main()
